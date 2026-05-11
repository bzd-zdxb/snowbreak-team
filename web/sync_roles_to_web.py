import json
from pathlib import Path
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent.parent   # demo 根目录
ROLES_DIR = BASE / "roles"
WEB_DIR = BASE / "web"
WEB_ROLES_DIR = WEB_DIR / "roles"
WEB_ROLES_DIR.mkdir(exist_ok=True)

# Copy images into web/roles
img_exts = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
img_map = {}
for p in ROLES_DIR.iterdir():
    if p.is_file() and p.suffix.lower() in img_exts:
        target = WEB_ROLES_DIR / p.name
        target.write_bytes(p.read_bytes())
        img_map[p.stem] = f"./roles/{p.name}"

xlsx_path = ROLES_DIR / "roles.xlsx"
wb = load_workbook(xlsx_path, data_only=True)
ws = wb.active

headers = {}
for c in range(1, 40):
    v = ws.cell(1, c).value
    if isinstance(v, str) and v.strip():
        headers[v.strip()] = c

required = ["name", "position", "damage_attr", "damage_type", "special_type", "version", "rank", "star"]
for k in required:
    headers.setdefault(k, None)

def norm_text(v):
    if v is None:
        return "待填写"
    s = str(v).strip()
    return s if s else "待填写"

def norm_int(v):
    try:
        return int(v)
    except Exception:
        return 0

roles = []
for r in range(2, ws.max_row + 1):
    name_col = headers["name"]
    if not name_col:
        continue
    name = ws.cell(r, name_col).value
    if not isinstance(name, str) or not name.strip():
        continue
    name = name.strip()

    row = {
        "name": name,
        "image": img_map.get(name, ""),
        "position": norm_text(ws.cell(r, headers["position"]).value) if headers["position"] else "待填写",
        "damage_attr": norm_text(ws.cell(r, headers["damage_attr"]).value) if headers["damage_attr"] else "待填写",
        "damage_type": norm_text(ws.cell(r, headers["damage_type"]).value) if headers["damage_type"] else "待填写",
        "special_type": norm_text(ws.cell(r, headers["special_type"]).value) if headers["special_type"] else "待填写",
        "version": norm_text(ws.cell(r, headers["version"]).value) if headers["version"] else "待填写",
        "rank": norm_int(ws.cell(r, headers["rank"]).value) if headers["rank"] else 0,
        "star": norm_int(ws.cell(r, headers["star"]).value) if headers["star"] else 0,
    }
    roles.append(row)

roles.sort(key=lambda x: x["name"].lower())
out = {"meta": {"source": "roles.xlsx", "count": len(roles)}, "roles": roles}
(WEB_DIR / "roles.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"synced roles.json, count={len(roles)}")
